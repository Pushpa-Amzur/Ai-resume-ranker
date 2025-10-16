"""
Resume Parser Module

This module provides functionality to parse different types of resume files
and extract text content for further processing.
"""

import os
import logging
from typing import Optional, Dict, Any
import PyPDF2
import pdfplumber
from docx import Document
import magic

logger = logging.getLogger(__name__)


class ResumeParser:
    """Parser for extracting text from resume files."""
    
    def __init__(self):
        self.supported_formats = ['.pdf', '.docx', '.txt']
    
    def parse_resume(self, file_path: str) -> Dict[str, Any]:
        """
        Parse a resume file and extract text content.
        
        Args:
            file_path (str): Path to the resume file
            
        Returns:
            Dict containing parsed resume data
        """
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"Resume file not found: {file_path}")
        
        file_extension = os.path.splitext(file_path)[1].lower()
        
        if file_extension not in self.supported_formats:
            raise ValueError(f"Unsupported file format: {file_extension}")
        
        try:
            if file_extension == '.pdf':
                text = self._parse_pdf(file_path)
            elif file_extension == '.docx':
                text = self._parse_docx(file_path)
            elif file_extension == '.txt':
                text = self._parse_txt(file_path)
            else:
                raise ValueError(f"Unsupported format: {file_extension}")
            
            return {
                'file_path': file_path,
                'file_name': os.path.basename(file_path),
                'file_extension': file_extension,
                'text_content': text,
                'word_count': len(text.split()),
                'character_count': len(text)
            }
            
        except Exception as e:
            logger.error(f"Error parsing resume {file_path}: {str(e)}")
            raise
    
    def _parse_pdf(self, file_path: str) -> str:
        """Extract text from PDF file using multiple methods."""
        text = ""
        
        # Try pdfplumber first (better for complex layouts)
        try:
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"
        except Exception as e:
            logger.warning(f"pdfplumber failed for {file_path}: {e}")
        
        # Fallback to PyPDF2 if pdfplumber didn't work
        if not text.strip():
            try:
                with open(file_path, 'rb') as file:
                    reader = PyPDF2.PdfReader(file)
                    for page in reader.pages:
                        page_text = page.extract_text()
                        if page_text:
                            text += page_text + "\n"
            except Exception as e:
                logger.error(f"PyPDF2 also failed for {file_path}: {e}")
                raise
        
        return text.strip()
    
    def _parse_docx(self, file_path: str) -> str:
        """Extract text from DOCX file."""
        try:
            doc = Document(file_path)
            text = ""
            
            # Extract text from paragraphs
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
            
            # Extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        text += cell.text + " "
                    text += "\n"
            
            return text.strip()
            
        except Exception as e:
            logger.error(f"Error parsing DOCX file {file_path}: {e}")
            raise
    
    def _parse_txt(self, file_path: str) -> str:
        """Extract text from TXT file."""
        try:
            with open(file_path, 'r', encoding='utf-8') as file:
                return file.read().strip()
        except UnicodeDecodeError:
            # Try with different encoding if UTF-8 fails
            try:
                with open(file_path, 'r', encoding='latin-1') as file:
                    return file.read().strip()
            except Exception as e:
                logger.error(f"Error parsing TXT file {file_path}: {e}")
                raise
        except Exception as e:
            logger.error(f"Error parsing TXT file {file_path}: {e}")
            raise
    
    def validate_file(self, file_path: str) -> bool:
        """
        Validate if a file can be processed.
        
        Args:
            file_path (str): Path to the file
            
        Returns:
            bool: True if file is valid and supported
        """
        if not os.path.exists(file_path):
            return False
        
        file_extension = os.path.splitext(file_path)[1].lower()
        
        if file_extension not in self.supported_formats:
            return False
        
        # Check file size (max 10MB)
        file_size = os.path.getsize(file_path)
        if file_size > 10 * 1024 * 1024:  # 10MB
            return False
        
        return True


def batch_parse_resumes(resume_directory: str) -> Dict[str, Dict[str, Any]]:
    """
    Parse multiple resume files from a directory.
    
    Args:
        resume_directory (str): Directory containing resume files
        
    Returns:
        Dict mapping file names to parsed resume data
    """
    parser = ResumeParser()
    results = {}
    
    if not os.path.exists(resume_directory):
        raise FileNotFoundError(f"Directory not found: {resume_directory}")
    
    for filename in os.listdir(resume_directory):
        file_path = os.path.join(resume_directory, filename)
        
        if os.path.isfile(file_path) and parser.validate_file(file_path):
            try:
                parsed_resume = parser.parse_resume(file_path)
                results[filename] = parsed_resume
                logger.info(f"Successfully parsed: {filename}")
            except Exception as e:
                logger.error(f"Failed to parse {filename}: {e}")
                results[filename] = {'error': str(e)}
    
    return results